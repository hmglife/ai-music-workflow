#!/usr/bin/env python3
"""Turn a designer's REAL requirement document into a reusable blank template.

This is how the skill preserves a designer's exact look-and-feel (their house font, row heights,
column widths, fills, borders): we take one of their real documents, clear the value column, tidy
the labels, and save it as a template. Filling that template later (build_requirement.py) inherits
all of that formatting for free.

Usage:
    python make_template.py --source <real_doc.xlsx> --out <template.xlsx> [--clear-values]

  --source        a real, representative requirement document
  --out           where to save the blank template
  --clear-values  replace column-B values with light placeholders (default: keep them as examples)

Label tidy-up: common inconsistent labels are normalized (e.g. "Music type" -> "Music Type",
"Requirement Description" -> "Requirement") while keeping each cell's original formatting.

Scope: operates on the workbook's ACTIVE sheet only (these briefs are single-sheet). If the source
has multiple sheets, the others are preserved in the saved file but only the active one is tidied;
a warning is printed so you can pick the right sheet if needed.

Dependency: openpyxl
"""
import sys
import os
import re as _re
import argparse

LABEL_FIX = {
    "music type": "Music Type",
    "requirement description": "Requirement",
    "requirement": "Requirement",
}

PLACEHOLDERS = {
    "Music Type": "<e.g. CG Music / Main Lobby Music / Story Music / Login Music>",
    "Duration": "<e.g. 35s / 1min>",
    "Order of Music": "<e.g. oneshot + intro/loop/end / loop>",
    "Format": "WAV 48KHz 24bit Stereo + Stem",
    "Name": "<project_music_<type>_<character>_<structure>>",
    "Sample Names": "<project_music_<type>_<character>_intro/loop/end>",
    "Reference Music": "<reference track name>",
    "Submit Time": "<YYYY-MM-DD>",
    "Requirement": "<requirement body>",
    "PS:": "<overall note: which sync points matter, how much freedom for the composer>",
}

# Column B on these rows is a SECTION HEADER, not a value — clearing it would destroy the
# document's structure (e.g. the CG timeline block is introduced by `Time | Requirements List`).
STRUCTURAL_LABELS = {"time"}

# Rows whose value is an embedded picture: they should end up genuinely empty, not placeholdered.
KEEP_EMPTY_LABELS = {"reference picture", "picture", "reference image"}

# Timeline segment rows (`0-Xs`, `X-Ys`, `10-24s` …). Matched by shape so a designer's own
# segment labels are covered too.
SEGMENT_RE = _re.compile(r"^\s*[\w.]+\s*-\s*[\w.]+s\s*$", _re.I)
SEGMENT_PLACEHOLDER = "<visuals in this segment + musical intent / freeze-frame or camera cut to align>"


def placeholder_for(key):
    """Decide what column B should become. Never silently blank an unknown label.

    An unknown label gets a generic `<...>` placeholder rather than "" for two reasons: the
    designer can see what the row is for, and `build_requirement.py --clean-placeholders` detects
    unused rows by exactly that `<...>` shape — an empty cell would be invisible to it.
    """
    low = key.lower()
    if low in STRUCTURAL_LABELS:
        return None                       # leave the header text untouched
    if low in KEEP_EMPTY_LABELS:
        return ""                         # picture row: genuinely empty
    if key in PLACEHOLDERS:
        return PLACEHOLDERS[key]
    if SEGMENT_RE.match(key):
        return SEGMENT_PLACEHOLDER
    return f"<{key.rstrip(':').strip().lower()}>"


def main():
    ap = argparse.ArgumentParser(add_help=True)
    ap.add_argument("--source", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--clear-values", action="store_true",
                    help="replace column-B values with placeholders instead of keeping examples")
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        here = os.path.dirname(os.path.abspath(__file__))
        print("[missing dependency] openpyxl is not available in this interpreter.\n"
              "  Install it automatically (do NOT ask the designer to run pip):\n"
              f'    "{sys.executable}" "{os.path.join(here, "ensure_deps.py")}" --for excel')
        sys.exit(2)

    wb = openpyxl.load_workbook(args.source)
    ws = wb.active
    if len(wb.worksheets) > 1:
        print(f"[warn] source has {len(wb.worksheets)} sheets; only the active sheet "
              f"'{ws.title}' is tidied (others kept as-is). "
              f"Sheets: {[s.title for s in wb.worksheets]}")

    generic = []
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1)
        if not a.value:
            continue
        raw = str(a.value).strip()
        key = raw.replace("\n", " ").strip()
        # normalize the label text (cell formatting preserved automatically)
        fixed = LABEL_FIX.get(key.lower())
        if fixed:
            a.value = fixed
            key = fixed
        if args.clear_values:
            ph = placeholder_for(key)
            if ph is None:
                continue                  # structural header — leave it alone
            b = ws.cell(row=r, column=2)
            if ph.startswith("<") and key not in PLACEHOLDERS and not SEGMENT_RE.match(key):
                generic.append(key)
            b.value = ph

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    wb.save(args.out)
    mode = "placeholders" if args.clear_values else "kept example values"
    print(f"[ok] template written: {args.out}  ({mode}; original formatting preserved)")
    if generic:
        print(f"[note] {len(generic)} label(s) had no built-in placeholder and got a generic one: "
              f"{generic}")
        print("       Open the template and reword those hints if you want them more helpful.")


if __name__ == "__main__":
    main()
